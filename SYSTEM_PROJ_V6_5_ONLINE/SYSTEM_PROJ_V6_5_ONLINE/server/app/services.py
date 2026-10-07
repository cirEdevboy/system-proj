from __future__ import annotations
from datetime import datetime, timedelta, timezone
from io import BytesIO
from pathlib import Path
import hashlib, os, re, smtplib, unicodedata, uuid
from email.message import EmailMessage
from difflib import SequenceMatcher
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from docx import Document as WordDocument
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
from sqlalchemy import func, or_, select
from sqlalchemy.orm import Session
from .config import settings
from .models import (
    AuditEvent, Document, Material, Notification, Project, ProjectItem,
    RomaneioItem, StageEvent, User
)
from .ws import manager

def utcnow():
    return datetime.now(timezone.utc)

def serialize_user(user: User):
    return {"id": user.id, "full_name": user.full_name, "email": user.email, "username": user.username, "badge": user.badge, "role": user.role, "status": user.status}

def serialize_project(p: Project):
    return {
        "id": p.id, "os": p.os, "client": p.client, "title": p.title,
        "status": p.status, "container_status": p.container_status,
        "transport_stage": p.transport_stage, "transfer_requested": p.transfer_requested,
        "collection_forecast": p.collection_forecast, "consumables_collected": p.consumables_collected,
        "version": p.version, "created_at": p.created_at.isoformat() if p.created_at else None,
        "updated_at": p.updated_at.isoformat() if p.updated_at else None,
        "initiated_by": serialize_user(p.initiator) if p.initiator else None,
        "current_responsible": serialize_user(p.responsible) if p.responsible else None,
    }

def audit(db: Session, actor: User | None, entity_type: str, entity_id, action: str, old=None, new=None, request=None):
    db.add(AuditEvent(
        actor_user_id=actor.id if actor else None,
        entity_type=entity_type,
        entity_id=str(entity_id),
        action=action,
        old_json=old,
        new_json=new,
        ip_address=request.client.host if request and request.client else None,
        user_agent=request.headers.get("user-agent") if request else None,
    ))

def create_notification(db: Session, user_id: int, kind: str, title: str, message: str, payload=None):
    n = Notification(user_id=user_id, kind=kind, title=title, message=message, payload_json=payload or {})
    db.add(n)
    db.flush()
    return n

async def push_notification(notification: Notification):
    await manager.send_user(notification.user_id, {
        "type": "notification",
        "notification": {
            "id": notification.id,
            "kind": notification.kind,
            "title": notification.title,
            "message": notification.message,
            "payload": notification.payload_json,
            "created_at": notification.created_at.isoformat() if notification.created_at else utcnow().isoformat(),
        }
    })

def normalize(text):
    text = str(text or "").strip().upper()
    text = unicodedata.normalize("NFD", text)
    text = "".join(c for c in text if unicodedata.category(c) != "Mn")
    return re.sub(r"\s+", " ", text)

def technical_numbers(text):
    return set(re.findall(r"\d+(?:[.,]\d+)?", normalize(text)))

def technical_words(text):
    stop = {"DE", "DA", "DO", "DAS", "DOS", "PARA", "COM", "E", "EM", "POR"}
    return {x for x in re.findall(r"[A-Z]+", normalize(text)) if len(x) >= 2 and x not in stop}

def score_material(desc, material_desc):
    a, b = normalize(desc), normalize(material_desc)
    if not a or not b: return 0.0
    if a == b: return 100.0
    na, nb = technical_numbers(a), technical_numbers(b)
    if na and nb and na.isdisjoint(nb): return 0.0
    wa, wb = technical_words(a), technical_words(b)
    if not wa or not wb: return SequenceMatcher(None, a, b).ratio() * 100
    inter = wa & wb
    coverage = len(inter)/max(len(wa),1)
    precision = len(inter)/max(len(wb),1)
    jaccard = len(inter)/max(len(wa|wb),1)
    seq = SequenceMatcher(None,a,b).ratio()
    score = coverage*45 + precision*20 + jaccard*15 + seq*20
    if na: score += 8 * len(na & nb)/len(na)
    return min(score,100)

UNITS = {"UN":"UN","UM":"UN","UND":"UN","M":"M","MT":"M","MTS":"M","METRO":"M","METROS":"M","M²":"M²","M2":"M²","LT":"LT","LTS":"LT","L":"LT","KG":"KG","CX":"CX","PAR":"PAR","PARES":"PAR"}

def norm_unit(value):
    return UNITS.get(str(value or "UN").strip().upper(), "UN")

def parse_word_bytes(data: bytes):
    doc = WordDocument(BytesIO(data))
    unit_regex = r"unidades|unidade|unid|und|um|un|mts²|mt²|m²|m2|metros|metro|mts|mt|m|litros|litro|lts|lt|l|quilos|quilo|kgs|kg|caixas|caixa|cxs|cx|pares|par|pcs|pçs|pc|pç|jg|kit|conj|cj"
    pattern = re.compile(rf"^\s*(?P<qtd>\d+(?:[.,]\d+)?)\s*(?P<unit>{unit_regex})?(?:\s*[-–—:]\s*|\s+)?(?P<desc>.+?)\s*$", re.I)
    code_re = re.compile(r"(?<!\d)(\d{8})(?!\d)")
    items=[]
    for p in doc.paragraphs:
        line = re.sub(r"^[\s•●▪■\-]+", "", p.text or "").strip()
        if not line: continue
        up = normalize(line)
        if any(x in up for x in ("LISTA DE CONSUMIVEIS","OS:","CLIENTE:","DATA:","REVISAO:")): continue
        code=""
        codes=code_re.findall(line)
        if codes:
            code=codes[-1]
            line=re.sub(rf"\s*[-–—]?\s*{re.escape(code)}\s*$", "", line).strip()
        m=pattern.match(line)
        qty=1; unit="UN"; desc=line
        if m:
            qty=float(m.group("qtd").replace(",","."))
            if qty.is_integer(): qty=int(qty)
            unit=norm_unit(m.group("unit")); desc=(m.group("desc") or "").strip(" -–—:")
        review = any(x in normalize(desc) for x in ("COMPRAR","ESPECIFICAR","EPIS PARA SOLDA","EPI PARA SOLDA"))
        if len(desc)>=3:
            items.append({"code":code,"description":desc,"quantity":qty,"unit":unit,"review":review})
    return items

def match_items(db: Session, items):
    materials = db.query(Material).filter(Material.active.is_(True)).all()
    by_code={m.code:m for m in materials}
    result=[]
    for item in items:
        match=None; score=0; state="NOT_FOUND"
        if item["code"] and item["code"] in by_code:
            match=by_code[item["code"]]; score=100; state="CODE"
        elif not item["review"]:
            best=None; best_score=0
            for m in materials:
                s=score_material(item["description"],m.description)
                if s>best_score: best_score=s; best=m
            score=best_score
            if best_score>=90: match=best; state="AUTO"
            elif best_score>=75: state="REVIEW"
        result.append((item,match,score,state))
    return result

def generate_transfer_xlsx(project: Project, items: list[ProjectItem], materials_by_id: dict[int, Material] | None = None):
    eligible=[i for i in items if i.transfer_eligible and i.material_code]
    if not eligible: raise ValueError("Não há itens transferíveis.")
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="TRANSFERENCIA"; ws.sheet_view.showGridLines=False
    blue="003B7A"; mid="0050A4"; light="EAF3FB"; white="FFFFFF"; border="D1D5DB"
    side=Side(style="thin",color=border); bd=Border(left=side,right=side,top=side,bottom=side)
    ws.merge_cells("A1:F1"); ws["A1"]="SOLICITAÇÃO DE TRANSFERÊNCIA DE MATERIAL"; ws["A1"].font=Font(size=18,bold=True,color=white); ws["A1"].fill=PatternFill("solid",fgColor=blue); ws["A1"].alignment=Alignment(horizontal="center")
    info=[("OS",project.os,"Cliente",project.client),("Projeto",project.title,"Responsável",project.responsible.full_name if project.responsible else "-")]
    r=3
    for a,b,c,d in info:
        ws.cell(r,1,a); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=3); ws.cell(r,2,b); ws.cell(r,4,c); ws.merge_cells(start_row=r,start_column=5,end_row=r,end_column=6); ws.cell(r,5,d)
        for col in range(1,7): ws.cell(r,col).border=bd
        ws.cell(r,1).fill=PatternFill("solid",fgColor=light); ws.cell(r,4).fill=PatternFill("solid",fgColor=light); ws.cell(r,1).font=Font(bold=True,color=blue); ws.cell(r,4).font=Font(bold=True,color=blue); r+=1
    headers=["Material","Nome do Material","Centro","Depósito","UM","Quantidade"]
    hr=7
    for c,h in enumerate(headers,1):
        cell=ws.cell(hr,c,h); cell.font=Font(bold=True,color=white); cell.fill=PatternFill("solid",fgColor=mid); cell.border=bd; cell.alignment=Alignment(horizontal="center")
    materials_by_id = materials_by_id or {}
    for idx,item in enumerate(eligible,hr+1):
        material_obj = materials_by_id.get(item.material_id) if item.material_id else None
        values=[
            item.material_code,
            item.matched_description or item.original_description,
            material_obj.center if material_obj else "",
            material_obj.deposit if material_obj else "",
            material_obj.unit if material_obj else item.unit,
            item.quantity,
        ]
        for c,v in enumerate(values,1):
            cell=ws.cell(idx,c,v); cell.border=bd; cell.alignment=Alignment(horizontal="left" if c==2 else "center",vertical="center",wrap_text=True)
    for c,w in enumerate([18,60,12,12,10,13],1): ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width=w
    ws.freeze_panes="A8"; ws.page_setup.orientation=ws.ORIENTATION_LANDSCAPE; ws.page_setup.fitToWidth=1; ws.sheet_properties.pageSetUpPr.fitToPage=True
    buf=BytesIO(); wb.save(buf); wb.close(); return buf.getvalue()

def generate_romaneio_xlsx(project: Project, items: list[RomaneioItem]):
    active=[x for x in items if not x.deleted_at]
    wb=openpyxl.Workbook(); ws=wb.active; ws.title="ROMANEIO"; ws.sheet_view.showGridLines=False
    blue="003B7A"; mid="0050A4"; white="FFFFFF"; light="EAF3FB"; border="D1D5DB"; side=Side(style="thin",color=border); bd=Border(left=side,right=side,top=side,bottom=side)
    ws.merge_cells("A1:H1"); ws["A1"]=f"ROMANEIO DE EMBARQUE - OS {project.os}"; ws["A1"].font=Font(size=18,bold=True,color=white); ws["A1"].fill=PatternFill("solid",fgColor=blue); ws["A1"].alignment=Alignment(horizontal="center")
    info=[("Cliente",project.client,"Projeto",project.title),("Responsável",project.responsible.full_name if project.responsible else "-","Emissão",utcnow().strftime("%d/%m/%Y %H:%M"))]
    r=3
    for a,b,c,d in info:
        ws.cell(r,1,a); ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4); ws.cell(r,2,b); ws.cell(r,5,c); ws.merge_cells(start_row=r,start_column=6,end_row=r,end_column=8); ws.cell(r,6,d)
        for col in range(1,9): ws.cell(r,col).border=bd
        for col in (1,5): ws.cell(r,col).fill=PatternFill("solid",fgColor=light); ws.cell(r,col).font=Font(bold=True,color=blue)
        r+=1
    headers=["SEQ","CATEGORIA","DESCRIÇÃO","QTD","UN","ORIGEM","DOCUMENTO ORIGEM","OBSERVAÇÃO"]
    hr=7
    for c,h in enumerate(headers,1):
        cell=ws.cell(hr,c,h); cell.font=Font(bold=True,color=white); cell.fill=PatternFill("solid",fgColor=mid); cell.border=bd; cell.alignment=Alignment(horizontal="center")
    for seq,item in enumerate(active,1):
        vals=[seq,item.category,item.description,item.quantity,item.unit,item.origin or "",item.source_document or "",item.observation or ""]
        rr=hr+seq
        for c,v in enumerate(vals,1):
            cell=ws.cell(rr,c,v); cell.border=bd; cell.alignment=Alignment(horizontal="left" if c in (2,3,6,7,8) else "center",vertical="center",wrap_text=True)
    for c,w in enumerate([7,18,58,10,8,18,30,34],1): ws.column_dimensions[openpyxl.utils.get_column_letter(c)].width=w
    ws.freeze_panes="A8"; ws.page_setup.orientation=ws.ORIENTATION_LANDSCAPE; ws.page_setup.fitToWidth=1; ws.sheet_properties.pageSetUpPr.fitToPage=True
    buf=BytesIO(); wb.save(buf); wb.close(); return buf.getvalue()

def generate_romaneio_pdf(project: Project, items: list[RomaneioItem]):
    active=[x for x in items if not x.deleted_at]
    buf=BytesIO(); doc=SimpleDocTemplate(buf,pagesize=landscape(A4),rightMargin=10*mm,leftMargin=10*mm,topMargin=10*mm,bottomMargin=10*mm)
    styles=getSampleStyleSheet(); story=[Paragraph(f"<b>ROMANEIO DE EMBARQUE - OS {project.os}</b>",styles["Title"]),Spacer(1,5*mm),Paragraph(f"Cliente: {project.client} &nbsp;&nbsp; Projeto: {project.title} &nbsp;&nbsp; Responsável: {project.responsible.full_name if project.responsible else '-'}",styles["Normal"]),Spacer(1,4*mm)]
    data=[["SEQ","CATEGORIA","DESCRIÇÃO","QTD","UN","ORIGEM","DOCUMENTO","OBSERVAÇÃO"]]
    for seq,x in enumerate(active,1): data.append([seq,x.category,x.description,x.quantity,x.unit,x.origin or "",x.source_document or "",x.observation or ""])
    table=Table(data,repeatRows=1,colWidths=[12*mm,30*mm,75*mm,15*mm,12*mm,30*mm,45*mm,45*mm])
    table.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#0050A4")),("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTNAME",(0,0),(-1,0),"Helvetica-Bold"),("GRID",(0,0),(-1,-1),0.4,colors.HexColor("#D1D5DB")),("VALIGN",(0,0),(-1,-1),"MIDDLE"),("FONTSIZE",(0,0),(-1,-1),7)]))
    story.append(table); doc.build(story); return buf.getvalue()

class Storage:
    def __init__(self):
        self.local = (
            settings.storage_mode.lower()
            != "s3"
        )

        self.client = None

        if self.local:
            Path(
                settings.local_storage_path
            ).mkdir(
                parents=True,
                exist_ok=True,
            )

        else:
            missing = [
                name
                for name, value in {
                    "S3_ENDPOINT": settings.s3_endpoint,
                    "S3_BUCKET": settings.s3_bucket,
                    "S3_ACCESS_KEY": settings.s3_access_key,
                    "S3_SECRET_KEY": settings.s3_secret_key,
                }.items()
                if not str(value or "").strip()
            ]

            if missing:
                raise RuntimeError(
                    "STORAGE_MODE=s3, mas faltam: "
                    + ", ".join(missing)
                )

            import boto3

            self.client = boto3.client(
                "s3",
                endpoint_url=settings.s3_endpoint,
                aws_access_key_id=settings.s3_access_key,
                aws_secret_access_key=settings.s3_secret_key,
                region_name=(
                    settings.s3_region
                    or "auto"
                ),
            )

            # No Cloudflare R2 o bucket deve ser criado previamente.
            # Não tentamos criá-lo automaticamente, porque tokens
            # restritos a um bucket podem não ter essa permissão.
            try:
                self.client.head_bucket(
                    Bucket=settings.s3_bucket
                )
            except Exception as exc:
                raise RuntimeError(
                    "Não foi possível acessar o bucket R2 configurado. "
                    "Confirme endpoint, bucket e credenciais."
                ) from exc
    def put(self,key,data,content_type="application/octet-stream"):
        if self.local:
            path=Path(settings.local_storage_path)/key; path.parent.mkdir(parents=True,exist_ok=True); path.write_bytes(data)
        else: self.client.put_object(Bucket=settings.s3_bucket,Key=key,Body=data,ContentType=content_type)
    def get(self,key):
        if self.local: return (Path(settings.local_storage_path)/key).read_bytes()
        return self.client.get_object(Bucket=settings.s3_bucket,Key=key)["Body"].read()
storage=Storage()

def save_versioned_document(db: Session, project_id:int, category:str, filename:str, data:bytes, mime:str, user_id:int, logical_key:str|None=None):
    logical_key=logical_key or re.sub(r"[^A-Za-z0-9_.-]+","_",filename).lower()
    latest=db.query(Document).filter(Document.project_id==project_id,Document.logical_key==logical_key,Document.deleted_at.is_(None)).order_by(Document.version.desc()).first()
    version=(latest.version+1) if latest else 1
    checksum=hashlib.sha256(data).hexdigest(); storage_key=f"projects/{project_id}/{logical_key}/v{version}_{uuid.uuid4().hex}_{filename}"
    storage.put(storage_key,data,mime)
    doc=Document(project_id=project_id,logical_key=logical_key,category=category,filename=filename,storage_key=storage_key,mime_type=mime,version=version,checksum_sha256=checksum,supersedes_id=latest.id if latest else None,uploaded_by_user_id=user_id)
    db.add(doc); db.flush(); return doc

def send_password_reset_email(email: str, token: str):
    if not settings.smtp_host: return False
    msg=EmailMessage(); msg["Subject"]="Project Center - Redefinição de senha"; msg["From"]=settings.smtp_from; msg["To"]=email
    link=f"{settings.public_base_url}/?reset_token={token}"
    msg.set_content(f"Use o link abaixo para redefinir sua senha. O link expira em 30 minutos.\n\n{link}")
    with smtplib.SMTP(settings.smtp_host,settings.smtp_port,timeout=20) as smtp:
        if settings.smtp_starttls: smtp.starttls()
        if settings.smtp_username: smtp.login(settings.smtp_username,settings.smtp_password)
        smtp.send_message(msg)
    return True
