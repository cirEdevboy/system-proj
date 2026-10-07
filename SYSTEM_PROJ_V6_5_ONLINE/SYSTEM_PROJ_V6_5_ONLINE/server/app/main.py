from __future__ import annotations
from datetime import datetime, timedelta, timezone
import asyncio
import time
from io import BytesIO
import os, secrets, uuid
import jwt
import openpyxl
from fastapi import FastAPI, Depends, File, Form, HTTPException, Request, UploadFile, WebSocket, WebSocketDisconnect, Query
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import HTMLResponse, Response
from fastapi.staticfiles import StaticFiles
from fastapi.templating import Jinja2Templates
from sqlalchemy import func, or_
from sqlalchemy.orm import Session
from .config import settings
from .database import Base, engine, get_db, SessionLocal
from .dependencies import get_current_user, require_admin, can_edit_project, require_project_editor
from .models import *
from .schemas import *
from .security import *
from .services import *
from .ws import manager
from .master_admin import ensure_master_admin, is_admin_role, is_master_admin
from .keepalive import keepalive_loop

Base.metadata.create_all(bind=engine)

# Garante a conta MASTER_ADMIN antes de aceitar logins.
with SessionLocal() as _bootstrap_db:
    ensure_master_admin(_bootstrap_db)

app=FastAPI(title="Project Center API",version="6.5.1",docs_url="/api/docs",openapi_url="/api/openapi.json")
_cors_origins = settings.cors_list or ["*"]
app.add_middleware(
    CORSMiddleware,
    allow_origins=_cors_origins,
    allow_credentials=("*" not in _cors_origins),
    allow_methods=["*"],
    allow_headers=["*"],
)
BASE=os.path.dirname(__file__); app.mount("/static",StaticFiles(directory=os.path.join(BASE,"static")),name="static"); templates=Jinja2Templates(directory=os.path.join(BASE,"templates"))

CONTAINER_STAGES=["AGUARDANDO DEFINIÇÃO DO CONTAINER","CONTAINER LOCADO","CONTAINER SENDO CARREGADO","CONTAINER ENVIADO","PROJETO SEM CONTAINER"]
TRANSPORT_STAGES=["AGUARDANDO PROGRAMAÇÃO DE TRANSPORTE","PROGRAMAÇÃO DE TRANSPORTE ALINHADA","CARREGANDO EQUIPAMENTOS NO CAMINHÃO","EQUIPAMENTOS ENVIADOS"]


@app.on_event("startup")
async def _start_keepalive():
    app.state.keepalive_task = None

    if settings.keepalive_enabled:
        app.state.keepalive_task = asyncio.create_task(
            keepalive_loop()
        )


@app.on_event("shutdown")
async def _stop_keepalive():
    task = getattr(
        app.state,
        "keepalive_task",
        None,
    )

    if task:
        task.cancel()

        try:
            await task
        except asyncio.CancelledError:
            pass


@app.get("/api/v1/health")
def health():
    return {
        "ok": True,
        "version": "6.5.1",
        "environment": settings.env,
        "storage_mode": settings.storage_mode,
        "database": (
            "sqlite"
            if settings.database_url.startswith("sqlite")
            else "postgresql"
        ),
        "keepalive": {
            "enabled": settings.keepalive_enabled,
            "interval_minutes": settings.keepalive_interval_minutes,
            "targets": len(settings.keepalive_url_list),
        },
    }

@app.get("/", response_class=HTMLResponse)
def web_home(request: Request):
    # Starlette/FastAPI atuais usam `request` como primeiro argumento
    # de TemplateResponse. O formato antigo ("index.html", contexto)
    # fazia o dicionário de contexto ser interpretado como nome do
    # template, causando: TypeError: unhashable type: 'dict'.
    return templates.TemplateResponse(
        request=request,
        name="index.html",
        context={
            "request": request,
            "app_name": settings.app_name,
        },
    )

@app.post("/api/v1/auth/register")
async def register(data:RegisterIn,request:Request,db:Session=Depends(get_db)):
    if db.query(User).filter(or_(func.lower(User.email)==data.email.lower(),func.lower(User.username)==data.username.lower(),User.badge==data.badge)).first():
        raise HTTPException(409,"E-mail, usuário ou crachá já cadastrado")
    user=User(full_name=data.full_name.strip(),email=data.email.lower(),username=data.username.lower().strip(),badge=data.badge.strip(),password_hash=hash_password(data.password),role="OPERATOR",status="PENDING")
    db.add(user); db.flush(); audit(db,user,"user",user.id,"REGISTER",new={"full_name":user.full_name,"email":user.email},request=request)
    notes=[]
    for admin in db.query(User).filter(User.role.in_(["ADMIN","MASTER_ADMIN"]),User.status=="ACTIVE",User.deleted_at.is_(None)).all():
        notes.append(create_notification(db,admin.id,"USER_APPROVAL","Novo usuário aguardando aprovação",f"{user.full_name} ({user.email}) solicitou acesso.",{"user_id":user.id}))
    db.commit()
    for n in notes: await push_notification(n)
    return {"message":"Cadastro realizado. Aguarde aprovação do administrador.","status":"PENDING"}


@app.post("/api/v1/auth/claim-migrated")
async def claim_migrated_account(
    data: MigratedAccountClaimIn,
    request: Request,
    db: Session = Depends(get_db),
):
    username = data.username.strip().lower()

    user = db.query(User).filter(
        func.lower(User.username) == username,
        User.deleted_at.is_(None),
    ).first()

    if not user or user.status != "MIGRATED":
        raise HTTPException(
            404,
            "Conta migrada não encontrada ou já ativada."
        )

    claim = db.query(MigrationClaim).filter(
        MigrationClaim.user_id == user.id,
        MigrationClaim.used_at.is_(None),
    ).first()

    if (
        not claim
        or claim.token_hash != sha256(
            data.activation_code.strip()
        )
    ):
        raise HTTPException(
            400,
            "Código de ativação inválido."
        )

    email = data.email.strip().lower()
    badge = data.badge.strip()

    email_owner = db.query(User).filter(
        func.lower(User.email) == email.lower(),
        User.id != user.id,
        User.deleted_at.is_(None),
    ).first()

    if email_owner:
        raise HTTPException(
            409,
            "Este e-mail já está em uso."
        )

    badge_owner = db.query(User).filter(
        User.badge == badge,
        User.id != user.id,
        User.deleted_at.is_(None),
    ).first()

    if badge_owner:
        raise HTTPException(
            409,
            "Este crachá já está em uso."
        )

    old = {
        "full_name": user.full_name,
        "email": user.email,
        "badge": user.badge,
        "status": user.status,
    }

    user.full_name = data.full_name.strip()
    user.email = email
    user.badge = badge
    user.password_hash = hash_password(data.password)
    user.status = "PENDING"

    claim.used_at = utcnow()

    audit(
        db,
        user,
        "user",
        user.id,
        "CLAIM_MIGRATED_ACCOUNT",
        old=old,
        new={
            "full_name": user.full_name,
            "email": user.email,
            "badge": user.badge,
            "status": "PENDING",
        },
        request=request,
    )

    for admin in db.query(User).filter(
        User.role.in_(["ADMIN", "MASTER_ADMIN"]),
        User.status == "ACTIVE",
        User.deleted_at.is_(None),
    ).all():
        create_notification(
            db,
            admin.id,
            "USER_PENDING_APPROVAL",
            "Conta migrada aguardando aprovação",
            (
                f"{user.full_name} completou os dados da conta "
                "migrada e aguarda aprovação."
            ),
            {"user_id": user.id},
        )

    db.commit()

    return {
        "message": (
            "Dados atualizados. Sua conta migrada agora aguarda "
            "aprovação do administrador."
        ),
        "status": "PENDING",
    }


@app.post("/api/v1/auth/login")
def login(data:LoginIn,request:Request,db:Session=Depends(get_db)):
    login=data.login.strip().lower(); user=db.query(User).filter(or_(func.lower(User.username)==login,func.lower(User.email)==login,User.badge==data.login.strip())).first()
    if not user or not verify_password(data.password,user.password_hash): raise HTTPException(401,"Credenciais inválidas")
    if user.status!="ACTIVE": raise HTTPException(403,f"Usuário com status {user.status}")
    sid=uuid.uuid4().hex; refresh=create_refresh_token(user.id,sid); access=create_access_token(user.id,sid)
    session=UserSession(id=sid,user_id=user.id,refresh_token_hash=sha256(refresh),device_name=data.device_name,ip_address=request.client.host if request.client else None,user_agent=request.headers.get("user-agent")); db.add(session); user.last_login_at=utcnow(); audit(db,user,"session",sid,"LOGIN",request=request); db.commit()
    return {"access_token":access,"refresh_token":refresh,"token_type":"bearer","user":serialize_user(user)}

@app.post("/api/v1/auth/refresh")
def refresh(data:RefreshIn,db:Session=Depends(get_db)):
    try: payload=decode_token(data.refresh_token,"refresh")
    except Exception: raise HTTPException(401,"Refresh token inválido")
    session=db.get(UserSession,payload.get("sid")); user=db.get(User,int(payload["sub"]))
    if not session or session.revoked_at or session.refresh_token_hash!=sha256(data.refresh_token) or not user or user.status!="ACTIVE": raise HTTPException(401,"Sessão inválida")
    access=create_access_token(user.id,session.id); return {"access_token":access,"token_type":"bearer"}

@app.post("/api/v1/auth/logout")
def logout(user:User=Depends(get_current_user),request:Request=None,db:Session=Depends(get_db)):
    token=(request.headers.get("authorization") or "").split(" ")[-1]
    try: sid=decode_token(token,"access").get("sid"); s=db.get(UserSession,sid); s.revoked_at=utcnow() if s else None; audit(db,user,"session",sid,"LOGOUT",request=request); db.commit()
    except Exception: pass
    return {"ok":True}

@app.get("/api/v1/auth/me")
def me(user:User=Depends(get_current_user)): return serialize_user(user)

@app.post("/api/v1/auth/request-password-reset")
async def request_reset(data:ResetRequestIn,db:Session=Depends(get_db)):
    user=db.query(User).filter(func.lower(User.email)==data.email.lower()).first()
    if user:
        raw=random_reset_token(); db.add(PasswordReset(user_id=user.id,token_hash=sha256(raw),expires_at=utcnow()+timedelta(minutes=30)))
        db.commit(); send_password_reset_email(user.email,raw)
    return {"message":"Se o e-mail existir, as instruções de redefinição serão enviadas."}

@app.post("/api/v1/auth/reset-password")
def reset_password(data:ResetPasswordIn,db:Session=Depends(get_db)):
    pr=db.query(PasswordReset).filter(PasswordReset.token_hash==sha256(data.token),PasswordReset.used_at.is_(None)).first()
    if not pr or pr.expires_at<utcnow(): raise HTTPException(400,"Token inválido ou expirado")
    user=db.get(User,pr.user_id); user.password_hash=hash_password(data.new_password); pr.used_at=utcnow();
    for s in db.query(UserSession).filter(UserSession.user_id==user.id,UserSession.revoked_at.is_(None)).all(): s.revoked_at=utcnow()
    db.commit(); return {"message":"Senha redefinida."}

@app.post("/api/v1/sessions/ping")
def ping(user:User=Depends(get_current_user)): return {"ok":True,"server_time":utcnow().isoformat()}

@app.get("/api/v1/sessions/my")
def my_sessions(user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    rows=db.query(UserSession).filter(UserSession.user_id==user.id).order_by(UserSession.created_at.desc()).all(); return [{"id":x.id,"device":x.device_name,"ip":x.ip_address,"last_seen":x.last_seen_at.isoformat(),"revoked":bool(x.revoked_at)} for x in rows]

@app.get("/api/v1/projects")
def list_projects(q:str="",page:int=1,page_size:int=50,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    qry=db.query(Project).filter(Project.deleted_at.is_(None)); q=q.strip()
    if q: qry=qry.filter(or_(Project.os.ilike(f"%{q}%"),Project.client.ilike(f"%{q}%"),Project.title.ilike(f"%{q}%")))
    total=qry.count(); rows=qry.order_by(Project.updated_at.desc()).offset((page-1)*page_size).limit(page_size).all(); return {"items":[serialize_project(x) for x in rows],"total":total,"page":page,"page_size":page_size}

@app.post("/api/v1/projects")
async def create_project(data:ProjectCreateIn,request:Request,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    if db.query(Project).filter(Project.os==data.os,Project.deleted_at.is_(None)).first(): raise HTTPException(409,"OS já cadastrada")
    p=Project(os=data.os.strip(),client=data.client.strip(),title=data.title.strip(),initiated_by_user_id=user.id,current_responsible_user_id=user.id,status="NOVO")
    db.add(p); db.flush(); db.add(ResponsibilityHistory(project_id=p.id,user_id=user.id,reason="Início do projeto")); db.add(StageEvent(project_id=p.id,area="project",stage="PROJETO INICIADO",user_id=user.id)); audit(db,user,"project",p.id,"CREATE",new=serialize_project(p),request=request); db.commit(); await manager.broadcast({"type":"project_update","project":serialize_project(p)}); return serialize_project(p)

@app.post("/api/v1/projects/from-word")
async def create_from_word(os_value:str=Form(...),client:str=Form(...),title:str=Form(...),file:UploadFile=File(...),request:Request=None,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    if db.query(Project).filter(Project.os==os_value,Project.deleted_at.is_(None)).first(): raise HTTPException(409,"OS já cadastrada")
    raw=await file.read(); parsed=parse_word_bytes(raw); matched=match_items(db,parsed)
    p=Project(os=os_value.strip(),client=client.strip(),title=title.strip(),initiated_by_user_id=user.id,current_responsible_user_id=user.id,status="PENDENTE COTAÇÃO" if any(not m for _,m,_,_ in matched) else "TRANSFERÊNCIA PENDENTE")
    db.add(p); db.flush(); db.add(ResponsibilityHistory(project_id=p.id,user_id=user.id,reason="Início do projeto")); db.add(StageEvent(project_id=p.id,area="project",stage="PROJETO INICIADO",user_id=user.id))
    for item,mat,score,state in matched:
        eligible = bool(mat and not item["review"])
        db.add(ProjectItem(project_id=p.id,original_description=item["description"],quantity=item["quantity"],unit=item["unit"],material_id=mat.id if mat else None,material_code=mat.code if mat else None,matched_description=mat.description if mat else None,match_score=round(score,2),match_state="REVIEW" if item["review"] else state,transfer_eligible=eligible))
        if eligible:
            db.add(RomaneioItem(
                project_id=p.id,
                category="CONSUMÍVEIS",
                description=mat.description,
                quantity=item["quantity"],
                unit=mat.unit or item["unit"],
                origin="TRANSFERENCIA",
                created_by_user_id=user.id,
            ))
    save_versioned_document(db,p.id,"LISTA_CONSUMIVEIS",file.filename,raw,file.content_type or "application/vnd.openxmlformats-officedocument.wordprocessingml.document",user.id,"lista_consumiveis")
    audit(db,user,"project",p.id,"CREATE_FROM_WORD",new={"items":len(parsed)},request=request); db.commit(); await manager.broadcast({"type":"project_update","project":serialize_project(p)}); return {"project":serialize_project(p),"items":len(parsed),"transferable":sum(1 for _,m,_,_ in matched if m)}

@app.get("/api/v1/projects/{project_id}")
def project_detail(project_id:int,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    p=db.get(Project,project_id)
    if not p or p.deleted_at: raise HTTPException(404,"Projeto não encontrado")
    result=serialize_project(p); result["can_edit"]=can_edit_project(p,user)
    result["stage_events"]=[{"area":e.area,"stage":e.stage,"occurred_at":e.occurred_at.isoformat(),"user_id":e.user_id} for e in db.query(StageEvent).filter(StageEvent.project_id==p.id).order_by(StageEvent.occurred_at.desc()).all()]
    result["responsibility_history"]=[{"user_id":h.user_id,"full_name":db.get(User,h.user_id).full_name if db.get(User,h.user_id) else "-","started_at":h.started_at.isoformat(),"ended_at":h.ended_at.isoformat() if h.ended_at else None,"reason":h.reason} for h in db.query(ResponsibilityHistory).filter(ResponsibilityHistory.project_id==p.id).order_by(ResponsibilityHistory.started_at).all()]
    return result

@app.put("/api/v1/projects/{project_id}/flow")
async def update_flow(project_id:int,data:FlowUpdateIn,request:Request,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    p=db.get(Project,project_id)
    if not p or p.deleted_at: raise HTTPException(404,"Projeto não encontrado")
    require_project_editor(p,user)
    if data.version!=p.version: raise HTTPException(409,"O projeto foi atualizado por outro usuário. Recarregue antes de salvar.")
    if data.container_status not in CONTAINER_STAGES or data.transport_stage not in TRANSPORT_STAGES: raise HTTPException(400,"Etapa inválida")
    old=serialize_project(p)
    changes=[]
    # A solicitação de transferência não é mais um checkbox manual.
    # Ela só muda para "solicitada" após o usuário confirmar que enviou
    # o rascunho preparado no Outlook.
    if p.consumables_collected!=data.consumables_collected:
        changes.append(("consumables","CONSUMÍVEIS COLETADOS" if data.consumables_collected else "AGUARDANDO COLETA"))
    if p.container_status!=data.container_status:
        changes.append(("container",data.container_status))
    if p.transport_stage!=data.transport_stage:
        changes.append(("transport",data.transport_stage))
    p.collection_forecast=data.collection_forecast
    p.consumables_collected=data.consumables_collected
    p.container_status=data.container_status
    p.transport_stage=data.transport_stage
    p.status=data.transport_stage if data.transport_stage!="AGUARDANDO PROGRAMAÇÃO DE TRANSPORTE" else data.container_status
    p.version+=1
    for area,stage in changes: db.add(StageEvent(project_id=p.id,area=area,stage=stage,user_id=user.id))
    audit(db,user,"project",p.id,"FLOW_UPDATE",old=old,new=serialize_project(p),request=request); db.commit(); await manager.broadcast({"type":"project_update","project":serialize_project(p)}); return serialize_project(p)

@app.post("/api/v1/projects/{project_id}/responsibility/request")
async def request_responsibility(project_id:int,data:ResponsibilityRequestIn,request:Request,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    p=db.get(Project,project_id)
    if not p or p.deleted_at: raise HTTPException(404,"Projeto não encontrado")
    require_project_editor(p,user); target=db.get(User,data.to_user_id)
    if not target or target.status!="ACTIVE": raise HTTPException(400,"Novo responsável não está ativo")
    if target.id==p.current_responsible_user_id: raise HTTPException(400,"Usuário já é responsável")
    if db.query(ResponsibilityTransfer).filter(ResponsibilityTransfer.project_id==p.id,ResponsibilityTransfer.status=="PENDING").first(): raise HTTPException(409,"Já existe solicitação pendente")
    tr=ResponsibilityTransfer(project_id=p.id,from_user_id=p.current_responsible_user_id,to_user_id=target.id,requested_by_user_id=user.id,reason=data.reason,status="PENDING"); db.add(tr); db.flush(); n=create_notification(db,target.id,"RESPONSIBILITY_TRANSFER",f"Solicitação para assumir OS {p.os}",f"{user.full_name} solicitou que você assuma o projeto {p.title}. Motivo: {data.reason}",{"transfer_id":tr.id,"project_id":p.id}); audit(db,user,"responsibility_transfer",tr.id,"REQUEST",new={"to_user_id":target.id,"reason":data.reason},request=request); db.commit(); await push_notification(n); return {"id":tr.id,"status":tr.status}

@app.post("/api/v1/responsibility-transfers/{transfer_id}/decision")
async def decide_transfer(transfer_id:int,data:ResponsibilityDecisionIn,request:Request,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    tr=db.get(ResponsibilityTransfer,transfer_id)
    if not tr or tr.status!="PENDING": raise HTTPException(404,"Solicitação não encontrada ou já respondida")
    if tr.to_user_id!=user.id: raise HTTPException(403,"Somente o destinatário pode responder")
    p=db.get(Project,tr.project_id); tr.responded_at=utcnow()
    if data.accept:
        old_resp=p.current_responsible_user_id; tr.status="ACCEPTED"; p.current_responsible_user_id=user.id; p.version+=1
        current=db.query(ResponsibilityHistory).filter(ResponsibilityHistory.project_id==p.id,ResponsibilityHistory.ended_at.is_(None)).order_by(ResponsibilityHistory.started_at.desc()).first()
        if current: current.ended_at=utcnow()
        db.add(ResponsibilityHistory(project_id=p.id,user_id=user.id,reason=tr.reason,transfer_id=tr.id)); db.add(StageEvent(project_id=p.id,area="responsibility",stage=f"RESPONSABILIDADE ASSUMIDA POR {user.full_name}",user_id=user.id))
        n=create_notification(db,old_resp,"RESPONSIBILITY_ACCEPTED",f"OS {p.os} transferida",f"{user.full_name} aceitou a responsabilidade pelo projeto.",{"project_id":p.id})
    else:
        tr.status="REJECTED"; n=create_notification(db,tr.from_user_id,"RESPONSIBILITY_REJECTED",f"Transferência da OS {p.os} recusada",f"{user.full_name} recusou a solicitação.",{"project_id":p.id})
    audit(db,user,"responsibility_transfer",tr.id,"ACCEPT" if data.accept else "REJECT",request=request); db.commit(); await push_notification(n); await manager.broadcast({"type":"project_update","project":serialize_project(p)}); return {"status":tr.status}

@app.post("/api/v1/projects/{project_id}/responsibility/admin-override")
async def admin_override(project_id:int,data:AdminResponsibilityOverrideIn,request:Request,admin:User=Depends(require_admin),db:Session=Depends(get_db)):
    p=db.get(Project,project_id); target=db.get(User,data.to_user_id)
    if not p or p.deleted_at or not target or target.status!="ACTIVE": raise HTTPException(400,"Projeto ou usuário inválido")
    old=p.current_responsible_user_id; tr=ResponsibilityTransfer(project_id=p.id,from_user_id=old,to_user_id=target.id,requested_by_user_id=admin.id,reason=data.reason,status="ACCEPTED",admin_override=True,responded_at=utcnow()); db.add(tr); db.flush()
    current=db.query(ResponsibilityHistory).filter(ResponsibilityHistory.project_id==p.id,ResponsibilityHistory.ended_at.is_(None)).first();
    if current: current.ended_at=utcnow()
    p.current_responsible_user_id=target.id; p.version+=1; db.add(ResponsibilityHistory(project_id=p.id,user_id=target.id,reason=f"Transferência administrativa: {data.reason}",transfer_id=tr.id)); db.add(StageEvent(project_id=p.id,area="responsibility",stage=f"RESPONSABILIDADE TRANSFERIDA POR ADMIN PARA {target.full_name}",user_id=admin.id)); n=create_notification(db,target.id,"RESPONSIBILITY_ADMIN",f"Você assumiu a OS {p.os}",f"ADMIN transferiu a responsabilidade para você. Motivo: {data.reason}",{"project_id":p.id}); audit(db,admin,"responsibility_transfer",tr.id,"ADMIN_OVERRIDE",old={"user_id":old},new={"user_id":target.id,"reason":data.reason},request=request); db.commit(); await push_notification(n); await manager.broadcast({"type":"project_update","project":serialize_project(p)}); return serialize_project(p)

@app.get("/api/v1/responsibility-transfers/inbox")
def transfer_inbox(user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    rows=db.query(ResponsibilityTransfer).filter(ResponsibilityTransfer.to_user_id==user.id,ResponsibilityTransfer.status=="PENDING").order_by(ResponsibilityTransfer.requested_at.desc()).all(); return [{"id":x.id,"project_id":x.project_id,"os":db.get(Project,x.project_id).os,"title":db.get(Project,x.project_id).title,"from_user":db.get(User,x.from_user_id).full_name,"reason":x.reason,"requested_at":x.requested_at.isoformat()} for x in rows]

@app.delete("/api/v1/projects/{project_id}")
async def soft_delete_project(project_id:int,request:Request,admin:User=Depends(require_admin),db:Session=Depends(get_db)):
    p=db.get(Project,project_id)
    if not p or p.deleted_at: raise HTTPException(404,"Projeto não encontrado")
    p.deleted_at=utcnow(); p.status="PROJETO EXCLUÍDO"; p.version+=1
    audit(db,admin,"project",p.id,"SOFT_DELETE",request=request); db.commit(); await manager.broadcast({"type":"project_deleted","project_id":p.id}); return {"ok":True}

@app.get("/api/v1/projects/{project_id}/items")
def project_items(project_id:int,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    return [{"id":x.id,"description":x.original_description,"quantity":x.quantity,"unit":x.unit,"material_code":x.material_code,"matched_description":x.matched_description,"score":x.match_score,"state":x.match_state,"transfer_eligible":x.transfer_eligible} for x in db.query(ProjectItem).filter(ProjectItem.project_id==project_id).all()]

TRANSFER_EMAIL_RECIPIENTS = [
    "nnicole@weg.net",
    "robertoa@weg.net",
    "luanderson@weg.net",
    "adrianosena@weg.net",
]


def _transfer_payload(project: Project, transferable: int):
    recipients = TRANSFER_EMAIL_RECIPIENTS
    subject = (
        f"Solicitação de Transferência | OS {project.os} | "
        f"{project.client} | {project.title}"
    )

    body_text = (
        "Bom dia,\n\n"
        "Segue em anexo a solicitação de transferência de materiais "
        "referente ao projeto abaixo:\n\n"
        f"OS: {project.os}\n"
        f"Cliente: {project.client}\n"
        f"Projeto: {project.title}\n"
        f"Quantidade de itens para transferência: {transferable}\n\n"
        "Favor prosseguir com a transferência dos materiais relacionados "
        "na planilha anexa.\n\n"
    )

    body_html = (
        "<p>Bom dia,</p>"
        "<p>Segue em anexo a solicitação de transferência de materiais "
        "referente ao projeto abaixo:</p>"
        "<table style='border-collapse:collapse'>"
        f"<tr><td><b>OS:</b></td><td style='padding-left:8px'>{project.os}</td></tr>"
        f"<tr><td><b>Cliente:</b></td><td style='padding-left:8px'>{project.client}</td></tr>"
        f"<tr><td><b>Projeto:</b></td><td style='padding-left:8px'>{project.title}</td></tr>"
        f"<tr><td><b>Itens:</b></td><td style='padding-left:8px'>{transferable}</td></tr>"
        "</table>"
        "<p>Favor prosseguir com a transferência dos materiais relacionados "
        "na planilha anexa.</p>"
    )

    return {
        "recipients": recipients,
        "subject": subject,
        "body_text": body_text,
        "body_html": body_html,
    }


@app.get("/api/v1/projects/{project_id}/transfer/status")
def transfer_status(
    project_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    p = db.get(Project, project_id)

    if not p or p.deleted_at:
        raise HTTPException(404, "Projeto não encontrado")

    items = db.query(ProjectItem).filter(
        ProjectItem.project_id == project_id
    ).all()

    transferable = sum(
        1
        for x in items
        if x.transfer_eligible and x.material_code
    )

    review = sum(
        1
        for x in items
        if not (x.transfer_eligible and x.material_code)
    )

    latest_doc = db.query(Document).filter(
        Document.project_id == project_id,
        Document.logical_key == "transferencia_solicitacao",
        Document.deleted_at.is_(None),
    ).order_by(
        Document.version.desc()
    ).first()

    last_generated = db.query(StageEvent).filter(
        StageEvent.project_id == project_id,
        StageEvent.area == "consumables",
        StageEvent.stage == "SOLICITAÇÃO DE TRANSFERÊNCIA GERADA",
    ).order_by(
        StageEvent.occurred_at.desc()
    ).first()

    last_sent = db.query(StageEvent).filter(
        StageEvent.project_id == project_id,
        StageEvent.area == "consumables",
        StageEvent.stage == "SOLICITAÇÃO DE TRANSFERÊNCIA ENVIADA",
    ).order_by(
        StageEvent.occurred_at.desc()
    ).first()

    return {
        "sent": bool(p.transfer_requested),
        "status": (
            "SOLICITAÇÃO ENVIADA"
            if p.transfer_requested
            else (
                "SOLICITAÇÃO GERADA"
                if latest_doc
                else "NÃO SOLICITADA"
            )
        ),
        "transferable": transferable,
        "review": review,
        "latest_document": (
            {
                "id": latest_doc.id,
                "filename": latest_doc.filename,
                "version": latest_doc.version,
                "created_at": latest_doc.created_at.isoformat(),
            }
            if latest_doc
            else None
        ),
        "generated_at": (
            last_generated.occurred_at.isoformat()
            if last_generated
            else None
        ),
        "sent_at": (
            last_sent.occurred_at.isoformat()
            if last_sent
            else None
        ),
    }


@app.post("/api/v1/projects/{project_id}/transfer/prepare")
async def prepare_transfer(
    project_id: int,
    request: Request,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    p = db.get(Project, project_id)

    if not p or p.deleted_at:
        raise HTTPException(404, "Projeto não encontrado")

    require_project_editor(p, user)

    items = db.query(ProjectItem).filter(
        ProjectItem.project_id == project_id
    ).all()

    eligible = [
        x
        for x in items
        if x.transfer_eligible and x.material_code
    ]

    if not eligible:
        raise HTTPException(
            400,
            "Não há itens transferíveis com código SAP para gerar a solicitação."
        )

    material_ids = [
        x.material_id
        for x in eligible
        if x.material_id
    ]

    mats = (
        db.query(Material).filter(
            Material.id.in_(material_ids)
        ).all()
        if material_ids
        else []
    )

    try:
        data = generate_transfer_xlsx(
            p,
            items,
            {m.id: m for m in mats},
        )
    except ValueError as exc:
        raise HTTPException(400, str(exc))

    filename = f"OS_{p.os}_TRANSFERENCIA.xlsx"

    doc = save_versioned_document(
        db,
        p.id,
        "TRANSFERENCIA_XLSX",
        filename,
        data,
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        user.id,
        logical_key="transferencia_solicitacao",
    )

    db.add(
        StageEvent(
            project_id=p.id,
            area="consumables",
            stage="SOLICITAÇÃO DE TRANSFERÊNCIA GERADA",
            user_id=user.id,
            metadata_json={
                "document_id": doc.id,
                "document_version": doc.version,
                "items": len(eligible),
            },
        )
    )

    audit(
        db,
        user,
        "project",
        p.id,
        "TRANSFER_PREPARED",
        new={
            "document_id": doc.id,
            "version": doc.version,
            "items": len(eligible),
        },
        request=request,
    )

    db.commit()

    mail = _transfer_payload(
        p,
        len(eligible),
    )

    return {
        "document_id": doc.id,
        "document_version": doc.version,
        "filename": filename,
        "transferable": len(eligible),
        "review": len(items) - len(eligible),
        "already_sent": bool(p.transfer_requested),
        **mail,
    }


@app.post("/api/v1/projects/{project_id}/transfer/confirm-sent")
async def confirm_transfer_sent(
    project_id: int,
    request: Request,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    p = db.get(Project, project_id)

    if not p or p.deleted_at:
        raise HTTPException(404, "Projeto não encontrado")

    require_project_editor(p, user)

    if p.transfer_requested:
        status = transfer_status(project_id, user, db)
        return {
            "ok": True,
            "already_confirmed": True,
            **status,
        }

    latest_doc = db.query(Document).filter(
        Document.project_id == project_id,
        Document.logical_key == "transferencia_solicitacao",
        Document.deleted_at.is_(None),
    ).order_by(
        Document.version.desc()
    ).first()

    if not latest_doc:
        raise HTTPException(
            400,
            "Gere a solicitação de transferência antes de confirmar o envio."
        )

    p.transfer_requested = True
    p.version += 1

    event = StageEvent(
        project_id=p.id,
        area="consumables",
        stage="SOLICITAÇÃO DE TRANSFERÊNCIA ENVIADA",
        user_id=user.id,
        metadata_json={
            "document_id": latest_doc.id,
            "document_version": latest_doc.version,
        },
    )

    db.add(event)

    audit(
        db,
        user,
        "project",
        p.id,
        "TRANSFER_CONFIRMED_SENT",
        old={"transfer_requested": False},
        new={
            "transfer_requested": True,
            "document_id": latest_doc.id,
            "document_version": latest_doc.version,
        },
        request=request,
    )

    db.commit()

    await manager.broadcast(
        {
            "type": "project_update",
            "project": serialize_project(p),
        }
    )

    return {
        "ok": True,
        "sent": True,
        "sent_at": event.occurred_at.isoformat(),
        "version": p.version,
    }


@app.get("/api/v1/projects/{project_id}/transfer.xlsx")
def transfer_xlsx(
    project_id: int,
    user: User = Depends(get_current_user),
    db: Session = Depends(get_db),
):
    p = db.get(Project, project_id)

    if not p or p.deleted_at:
        raise HTTPException(404, "Projeto não encontrado")

    items = db.query(ProjectItem).filter(
        ProjectItem.project_id == project_id
    ).all()

    material_ids = [
        x.material_id
        for x in items
        if x.material_id
    ]

    mats = (
        db.query(Material).filter(
            Material.id.in_(material_ids)
        ).all()
        if material_ids
        else []
    )

    try:
        data = generate_transfer_xlsx(
            p,
            items,
            {m.id: m for m in mats},
        )
    except ValueError as exc:
        raise HTTPException(400, str(exc))

    return Response(
        data,
        media_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
        headers={
            "Content-Disposition": (
                f'attachment; filename="OS_{p.os}_TRANSFERENCIA.xlsx"'
            )
        },
    )

@app.get("/api/v1/projects/{project_id}/romaneio")
def romaneio_list(project_id:int,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    rows=db.query(RomaneioItem).filter(RomaneioItem.project_id==project_id,RomaneioItem.deleted_at.is_(None)).order_by(RomaneioItem.category,RomaneioItem.id).all(); return [{"id":x.id,"category":x.category,"description":x.description,"quantity":x.quantity,"unit":x.unit,"origin":x.origin,"source_document":x.source_document,"observation":x.observation} for x in rows]

@app.post("/api/v1/projects/{project_id}/romaneio")
async def romaneio_add(project_id:int,data:RomaneioItemIn,request:Request,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    p=db.get(Project,project_id); require_project_editor(p,user); item=RomaneioItem(project_id=p.id,category=data.category.upper(),description=data.description,quantity=data.quantity,unit=data.unit.upper(),origin="MANUAL",source_document=data.source_document,observation=data.observation,created_by_user_id=user.id); db.add(item); audit(db,user,"romaneio_item","new","CREATE",new=data.model_dump(),request=request); db.commit(); await manager.broadcast({"type":"project_update","project":serialize_project(p)}); return {"id":item.id}

@app.get("/api/v1/projects/{project_id}/romaneio.xlsx")
def romaneio_xlsx(project_id:int,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    p=db.get(Project,project_id); items=db.query(RomaneioItem).filter(RomaneioItem.project_id==project_id).all(); data=generate_romaneio_xlsx(p,items); return Response(data,media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="OS_{p.os}_ROMANEIO.xlsx"'})

@app.get("/api/v1/projects/{project_id}/romaneio.pdf")
def romaneio_pdf(project_id:int,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    p=db.get(Project,project_id); items=db.query(RomaneioItem).filter(RomaneioItem.project_id==project_id).all(); data=generate_romaneio_pdf(p,items); return Response(data,media_type="application/pdf",headers={"Content-Disposition":f'attachment; filename="OS_{p.os}_ROMANEIO.pdf"'})

@app.post("/api/v1/projects/{project_id}/documents")
async def upload_document(project_id:int,category:str=Form(...),logical_key:str=Form(""),file:UploadFile=File(...),request:Request=None,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    p=db.get(Project,project_id); require_project_editor(p,user); data=await file.read(); doc=save_versioned_document(db,p.id,category,file.filename,data,file.content_type or "application/octet-stream",user.id,logical_key or None); audit(db,user,"document",doc.id,"UPLOAD",new={"category":category,"filename":file.filename,"version":doc.version},request=request); db.commit(); return {"id":doc.id,"version":doc.version}

@app.get("/api/v1/projects/{project_id}/documents")
def list_documents(project_id:int,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    rows=db.query(Document).filter(Document.project_id==project_id,Document.deleted_at.is_(None)).order_by(Document.logical_key,Document.version.desc()).all(); return [{"id":x.id,"logical_key":x.logical_key,"category":x.category,"filename":x.filename,"version":x.version,"created_at":x.created_at.isoformat()} for x in rows]

@app.get("/api/v1/documents/{document_id}/download")
def download_document(document_id:int,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    doc=db.get(Document,document_id)
    if not doc or doc.deleted_at: raise HTTPException(404,"Documento não encontrado")
    data=storage.get(doc.storage_key); return Response(data,media_type=doc.mime_type or "application/octet-stream",headers={"Content-Disposition":f'attachment; filename="{doc.filename}"'})

@app.delete("/api/v1/documents/{document_id}")
def soft_delete_document(document_id:int,request:Request,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    doc=db.get(Document,document_id)
    if not doc or doc.deleted_at: raise HTTPException(404,"Documento não encontrado")
    p=db.get(Project,doc.project_id); require_project_editor(p,user)
    doc.deleted_at=utcnow(); audit(db,user,"document",doc.id,"SOFT_DELETE",request=request); db.commit(); return {"ok":True}

@app.get("/api/v1/materials")
def materials(q:str="",page:int=1,page_size:int=100,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    qry=db.query(Material); q=q.strip()
    if q: qry=qry.filter(or_(Material.code.ilike(f"%{q}%"),Material.description.ilike(f"%{q}%"),Material.center.ilike(f"%{q}%"),Material.deposit.ilike(f"%{q}%")))
    total=qry.count(); rows=qry.order_by(Material.active.desc(),Material.description).offset((page-1)*page_size).limit(page_size).all(); return {"items":[{"id":x.id,"code":x.code,"description":x.description,"center":x.center,"deposit":x.deposit,"unit":x.unit,"active":x.active} for x in rows],"total":total,"page":page,"page_size":page_size}

@app.post("/api/v1/materials/import")
async def import_materials(
    file: UploadFile = File(...),
    replace: bool = Form(False),
    admin: User = Depends(require_admin),
    db: Session = Depends(get_db),
):
    """
    Importa o catálogo de insumos via XLSX/XLSM.

    V6.5.1:
    - lê o arquivo em modo read_only;
    - identifica a planilha CONSULTA automaticamente;
    - deduplica por (código, centro, depósito) antes de acessar o banco;
    - carrega o catálogo existente em UMA consulta;
    - elimina milhares de SELECTs individuais no PostgreSQL;
    - retorna diagnóstico completo para a interface.
    """
    started = time.perf_counter()

    filename = (file.filename or "").strip()
    suffix = os.path.splitext(filename.lower())[1]

    if suffix not in {".xlsx", ".xlsm"}:
        raise HTTPException(
            400,
            "Formato inválido. Selecione um arquivo .xlsx ou .xlsm."
        )

    raw = await file.read()

    if not raw:
        raise HTTPException(400, "O arquivo enviado está vazio.")

    # Limite folgado para evitar upload acidental de arquivos enormes
    # no Web Service gratuito.
    if len(raw) > 25 * 1024 * 1024:
        raise HTTPException(
            400,
            "Arquivo muito grande. O limite do catálogo é 25 MB."
        )

    try:
        wb = openpyxl.load_workbook(
            BytesIO(raw),
            data_only=True,
            read_only=True,
        )
    except Exception as exc:
        raise HTTPException(
            400,
            f"Não foi possível abrir o Excel: {exc}"
        )

    try:
        ws = (
            wb["CONSULTA"]
            if "CONSULTA" in wb.sheetnames
            else wb.active
        )

        aliases = {
            "code": {
                "MATERIAL",
                "CODIGO",
                "CÓDIGO",
            },
            "description": {
                "TEXTO BREVE DE MATERIAL",
                "DESCRICAO",
                "DESCRIÇÃO",
                "NOME DO MATERIAL",
            },
            "center": {
                "CENTRO",
            },
            "deposit": {
                "DEPOSITO",
                "DEPÓSITO",
            },
            "unit": {
                "UM BASICA",
                "UM BÁSICA",
                "UNIDADE",
                "UM",
            },
        }

        normalized_aliases = {
            key: {
                normalize(value)
                for value in values
            }
            for key, values in aliases.items()
        }

        headers = {}
        header_row = 1

        for row_index in range(
            1,
            min(ws.max_row, 20) + 1,
        ):
            found = {}

            for col_index in range(
                1,
                min(ws.max_column, 20) + 1,
            ):
                header = normalize(
                    ws.cell(
                        row_index,
                        col_index,
                    ).value
                )

                for key, values in normalized_aliases.items():
                    if header in values:
                        found[key] = col_index

            if (
                "code" in found
                and "description" in found
            ):
                headers = found
                header_row = row_index
                break

        # Compatibilidade com o arquivo histórico:
        # Id_Inicial | Material | Texto breve | Centro | Depósito | UM
        if (
            not headers
            and ws.max_column >= 6
        ):
            headers = {
                "code": 2,
                "description": 3,
                "center": 4,
                "deposit": 5,
                "unit": 6,
            }
            header_row = 1

        if (
            "code" not in headers
            or "description" not in headers
        ):
            raise HTTPException(
                400,
                "Não encontrei as colunas Material/Código e "
                "Texto breve de material/Descrição."
            )

        records = {}
        ignored = 0
        duplicated_in_file = 0

        for row in ws.iter_rows(
            min_row=header_row + 1,
            values_only=True,
        ):
            def val(key):
                index = headers.get(key)

                if (
                    not index
                    or index > len(row)
                ):
                    return None

                return row[index - 1]

            code = str(
                val("code") or ""
            ).strip()

            description = str(
                val("description") or ""
            ).strip()

            center = str(
                val("center") or ""
            ).strip()

            deposit = str(
                val("deposit") or ""
            ).strip()

            unit = norm_unit(
                val("unit")
            )

            if (
                code.endswith(".0")
                and code[:-2].isdigit()
            ):
                code = code[:-2]

            if (
                not code
                or not description
            ):
                ignored += 1
                continue

            key = (
                code,
                center,
                deposit,
            )

            if key in records:
                duplicated_in_file += 1

            records[key] = {
                "code": code,
                "description": description,
                "center": center,
                "deposit": deposit,
                "unit": unit,
            }

    finally:
        wb.close()

    if not records:
        raise HTTPException(
            400,
            "Nenhum insumo válido foi encontrado no arquivo."
        )

    # UMA consulta em vez de um SELECT por linha.
    existing_rows = db.query(Material).all()

    existing = {
        (
            item.code,
            item.center,
            item.deposit,
        ): item
        for item in existing_rows
    }

    if replace:
        db.query(Material).update(
            {
                Material.active: False
            },
            synchronize_session=False,
        )

    created = 0
    updated = 0
    unchanged = 0

    new_objects = []

    for key, row in records.items():
        material = existing.get(key)

        if material is None:
            new_objects.append(
                Material(
                    code=row["code"],
                    center=row["center"],
                    deposit=row["deposit"],
                    description=row["description"],
                    unit=row["unit"],
                    active=True,
                )
            )
            created += 1
            continue

        changed = (
            material.description != row["description"]
            or material.unit != row["unit"]
            or not material.active
        )

        material.description = row["description"]
        material.unit = row["unit"]
        material.active = True

        if changed:
            updated += 1
        else:
            unchanged += 1

    if new_objects:
        db.add_all(new_objects)

    audit(
        db,
        admin,
        "material_catalog",
        "all",
        "IMPORT",
        new={
            "filename": filename,
            "sheet": ws.title if "ws" in locals() else None,
            "records": len(records),
            "created": created,
            "updated": updated,
            "unchanged": unchanged,
            "ignored": ignored,
            "duplicates_in_file": duplicated_in_file,
            "replace": replace,
        },
    )

    db.commit()

    active_total = db.query(Material).filter(
        Material.active.is_(True)
    ).count()

    elapsed = round(
        time.perf_counter() - started,
        2,
    )

    payload = {
        "processed": len(records),
        "created": created,
        "updated": updated,
        "unchanged": unchanged,
        "ignored": ignored,
        "duplicates_in_file": duplicated_in_file,
        "active_total": active_total,
        "sheet": ws.title if "ws" in locals() else None,
        "header_row": header_row,
        "elapsed_seconds": elapsed,
    }

    await manager.broadcast(
        {
            "type": "catalog_update",
            **payload,
        }
    )

    return payload

@app.get("/api/v1/notifications")
def notifications(user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    rows=db.query(Notification).filter(Notification.user_id==user.id).order_by(Notification.created_at.desc()).limit(100).all(); return [{"id":x.id,"kind":x.kind,"title":x.title,"message":x.message,"payload":x.payload_json,"read":bool(x.read_at),"created_at":x.created_at.isoformat()} for x in rows]

@app.post("/api/v1/notifications/{notification_id}/read")
def mark_read(notification_id:int,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    n=db.get(Notification,notification_id)
    if not n or n.user_id!=user.id: raise HTTPException(404,"Notificação não encontrada")
    n.read_at=utcnow(); db.commit(); return {"ok":True}

@app.get("/api/v1/search")
def global_search(q:str,user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    q=q.strip();
    if len(q)<2: return {"projects":[],"materials":[],"users":[],"documents":[]}
    projects=db.query(Project).filter(Project.deleted_at.is_(None),or_(Project.os.ilike(f"%{q}%"),Project.client.ilike(f"%{q}%"),Project.title.ilike(f"%{q}%"))).limit(20).all()
    mats=db.query(Material).filter(or_(Material.code.ilike(f"%{q}%"),Material.description.ilike(f"%{q}%"))).limit(20).all()
    users=db.query(User).filter(User.deleted_at.is_(None),or_(User.full_name.ilike(f"%{q}%"),User.email.ilike(f"%{q}%"),User.badge.ilike(f"%{q}%"))).limit(20).all()
    docs=db.query(Document).filter(Document.deleted_at.is_(None),or_(Document.filename.ilike(f"%{q}%"),Document.category.ilike(f"%{q}%"))).limit(20).all()
    return {"projects":[serialize_project(x) for x in projects],"materials":[{"id":x.id,"code":x.code,"description":x.description} for x in mats],"users":[serialize_user(x) for x in users],"documents":[{"id":x.id,"project_id":x.project_id,"filename":x.filename,"category":x.category,"version":x.version} for x in docs]}

@app.get("/api/v1/admin/users")
def admin_users(status_filter:str="",admin:User=Depends(require_admin),db:Session=Depends(get_db)):
    q=db.query(User).filter(User.deleted_at.is_(None));
    if status_filter: q=q.filter(User.status==status_filter)
    return [serialize_user(x) for x in q.order_by(User.created_at.desc()).all()]

@app.post("/api/v1/admin/users/{user_id}/approve")
async def approve_user(
    user_id:int,
    data:UserApproveIn,
    request:Request,
    admin:User=Depends(require_admin),
    db:Session=Depends(get_db)
):
    user=db.get(User,user_id)

    if not user:
        raise HTTPException(404,"Usuário não encontrado")

    if is_master_admin(user):
        raise HTTPException(
            403,
            "A conta MASTER_ADMIN é protegida e não pode ser alterada."
        )

    requested_role=data.role.upper().strip()
    allowed={"OPERATOR","CONSULTA","ADMIN"}

    if requested_role not in allowed:
        raise HTTPException(400,"Perfil inválido")

    # Somente o MASTER_ADMIN pode criar/promover outros administradores.
    if requested_role=="ADMIN" and not is_master_admin(admin):
        raise HTTPException(
            403,
            "Somente o MASTER_ADMIN pode conceder perfil ADMIN."
        )

    user.status="ACTIVE"
    user.role=requested_role
    user.approved_at=utcnow()
    user.approved_by_id=admin.id

    n=create_notification(
        db,
        user.id,
        "ACCOUNT_APPROVED",
        "Acesso aprovado",
        "Seu cadastro no Project Center foi aprovado.",
        {}
    )

    audit(
        db,
        admin,
        "user",
        user.id,
        "APPROVE",
        new={"role":user.role},
        request=request
    )

    db.commit()
    await push_notification(n)
    return serialize_user(user)

@app.post("/api/v1/admin/users/{user_id}/block")
def block_user(
    user_id:int,
    request:Request,
    admin:User=Depends(require_admin),
    db:Session=Depends(get_db)
):
    user=db.get(User,user_id)

    if not user:
        raise HTTPException(404,"Usuário não encontrado")

    if is_master_admin(user):
        raise HTTPException(
            403,
            "A conta MASTER_ADMIN é protegida e não pode ser bloqueada."
        )

    # ADMIN comum não bloqueia outro ADMIN.
    if user.role=="ADMIN" and not is_master_admin(admin):
        raise HTTPException(
            403,
            "Somente o MASTER_ADMIN pode bloquear outro ADMIN."
        )

    user.status="BLOCKED"

    for s in db.query(UserSession).filter(
        UserSession.user_id==user.id,
        UserSession.revoked_at.is_(None)
    ).all():
        s.revoked_at=utcnow()

    audit(
        db,
        admin,
        "user",
        user.id,
        "BLOCK",
        request=request
    )

    db.commit()
    return serialize_user(user)

@app.get("/api/v1/admin/online-users")
def online_users(admin:User=Depends(require_admin),db:Session=Depends(get_db)):
    since=utcnow()-timedelta(seconds=90); rows=db.query(UserSession).filter(UserSession.revoked_at.is_(None),UserSession.last_seen_at>=since).order_by(UserSession.last_seen_at.desc()).all(); return [{"session_id":s.id,"user":serialize_user(db.get(User,s.user_id)),"device":s.device_name,"ip":s.ip_address,"last_seen":s.last_seen_at.isoformat()} for s in rows]

@app.post("/api/v1/admin/sessions/{session_id}/revoke")
def revoke_session(session_id:str,request:Request,admin:User=Depends(require_admin),db:Session=Depends(get_db)):
    s=db.get(UserSession,session_id)
    if not s: raise HTTPException(404,"Sessão não encontrada")
    s.revoked_at=utcnow(); audit(db,admin,"session",session_id,"REVOKE",request=request); db.commit(); return {"ok":True}

@app.get("/api/v1/admin/activity")
def activity(limit:int=100,admin:User=Depends(require_admin),db:Session=Depends(get_db)):
    rows=db.query(AuditEvent).order_by(AuditEvent.created_at.desc()).limit(min(limit,500)).all(); return [{"id":x.id,"actor":db.get(User,x.actor_user_id).full_name if x.actor_user_id and db.get(User,x.actor_user_id) else "SYSTEM","entity_type":x.entity_type,"entity_id":x.entity_id,"action":x.action,"old":x.old_json,"new":x.new_json,"ip":x.ip_address,"created_at":x.created_at.isoformat()} for x in rows]

@app.get("/api/v1/users/active")
def active_users(user:User=Depends(get_current_user),db:Session=Depends(get_db)):
    return [serialize_user(x) for x in db.query(User).filter(User.status=="ACTIVE",User.deleted_at.is_(None)).order_by(User.full_name).all()]

@app.websocket("/api/v1/ws")
async def websocket_endpoint(websocket:WebSocket,token:str=Query(...)):
    db=next(get_db())
    try:
        payload=decode_token(token,"access"); user=db.get(User,int(payload["sub"])); session=db.get(UserSession,payload.get("sid"))
        if not user or user.status!="ACTIVE" or not session or session.revoked_at: await websocket.close(code=4401); return
        await manager.connect(user.id,websocket)
        while True:
            msg=await websocket.receive_text()
            if msg=="ping": await websocket.send_json({"type":"pong","time":utcnow().isoformat()})
    except WebSocketDisconnect: pass
    except Exception:
        try: await websocket.close(code=4401)
        except Exception: pass
    finally:
        try: await manager.disconnect(user.id,websocket)
        except Exception: pass
        db.close()
